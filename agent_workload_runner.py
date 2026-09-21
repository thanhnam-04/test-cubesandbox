"""Bounded, reproducible agent-like jobs. Runs as a child inside /scratch."""

import csv
import hashlib
import json
import os
import random
import subprocess
import sys
import time
from collections import Counter
from pathlib import Path


WORKLOADS = ("documents", "data", "code", "images", "search", "workflow")


def make_data(directory, rows=30000):
    path = directory / "sales.csv"
    rng = random.Random(42)
    totals = Counter()
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["day", "region", "product", "amount"])
        for number in range(rows):
            region = ("north", "south", "east", "west")[number % 4]
            amount = rng.randrange(10, 500)
            writer.writerow([number % 365, region, number % 17, amount])
            totals[region] += amount
    observed = Counter()
    with path.open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        for row in reader:
            observed[row["region"]] += int(row["amount"])
    if observed != totals:
        raise RuntimeError("CSV aggregate validation failed")
    os.environ["MPLCONFIGDIR"] = str(directory / ".mplconfig")
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib import pyplot as plt
    figure, axes = plt.subplots(figsize=(7, 4))
    axes.bar(list(observed), list(observed.values()))
    axes.set_title("Revenue by region")
    figure.tight_layout()
    chart = directory / "chart.png"
    figure.savefig(chart, dpi=110)
    plt.close(figure)
    if chart.stat().st_size < 1000:
        raise RuntimeError("chart is empty")
    return {"rows": rows, "regions": dict(observed), "chart_bytes": chart.stat().st_size}


def make_documents(directory, chart=None):
    from docx import Document
    from pypdf import PdfReader
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    doc = Document()
    doc.add_heading("Agent workload report", 0)
    for number in range(120):
        doc.add_paragraph(f"Finding {number + 1}: " + "Validated sample data. " * 8)
    docx_path = directory / "report.docx"
    doc.save(docx_path)
    parsed = Document(docx_path)
    if len(parsed.paragraphs) != 121:
        raise RuntimeError("DOCX paragraph validation failed")
    pdf_path = directory / "report.pdf"
    pdf = canvas.Canvas(str(pdf_path))
    for page in range(4):
        pdf.setFont("Helvetica", 13)
        pdf.drawString(48, 790, f"Agent report - page {page + 1}")
        for line in range(35):
            pdf.drawString(48, 760 - line * 18, f"Finding {page * 35 + line + 1}: validated sample data")
        if chart is not None and page == 0:
            pdf.drawImage(ImageReader(str(chart)), 300, 520, width=250, height=145)
        pdf.showPage()
    pdf.save()
    if len(PdfReader(str(pdf_path)).pages) != 4:
        raise RuntimeError("PDF page validation failed")
    return {"docx_bytes": docx_path.stat().st_size,
            "pdf_bytes": pdf_path.stat().st_size, "pdf_pages": 4}


def run_code(directory):
    (directory / "calculator.py").write_text(
        "def prime(n):\n"
        "    if n < 2: return False\n"
        "    for d in range(2, int(n ** 0.5) + 1):\n"
        "        if n % d == 0: return False\n"
        "    return True\n"
        "def count_primes(limit):\n"
        "    return sum(prime(n) for n in range(limit))\n", encoding="utf-8")
    (directory / "test_calculator.py").write_text(
        "import unittest\nfrom calculator import count_primes, prime\n"
        "class CalculatorTest(unittest.TestCase):\n"
        "    def test_prime(self): self.assertTrue(prime(7919))\n"
        "    def test_composite(self): self.assertFalse(prime(8000))\n"
        "    def test_count(self): self.assertEqual(count_primes(10000), 1229)\n"
        "if __name__ == '__main__': unittest.main()\n", encoding="utf-8")
    result = subprocess.run([sys.executable, "-m", "unittest", "discover", "-s", str(directory)],
                            cwd=directory, capture_output=True, text=True, timeout=20, check=False)
    if result.returncode != 0 or "Ran 3 tests" not in result.stderr:
        raise RuntimeError("generated code tests failed: " + result.stderr[-300:])
    return {"tests_passed": 3, "test_output": result.stderr[-120:]}


def make_images(directory):
    from PIL import Image, ImageDraw
    widths = []
    thumbnails = []
    for index in range(4):
        photo = Image.new("RGB", (1600, 1000), (35 + index * 40, 70, 145))
        pen = ImageDraw.Draw(photo)
        for line in range(0, 1000, 20):
            pen.line((0, line, 1599, (line + index * 21) % 1000), fill=(220, 190, 80), width=3)
        source = directory / f"source-{index}.png"
        photo.save(source)
        with Image.open(source) as loaded:
            thumb = loaded.resize((400, 250))
        thumb.save(directory / f"thumb-{index}.jpg", quality=85)
        thumbnails.append(thumb)
        widths.append(thumb.width)
    contact = Image.new("RGB", (800, 500))
    for index, thumb in enumerate(thumbnails):
        contact.paste(thumb, ((index % 2) * 400, (index // 2) * 250))
    contact_path = directory / "contact.jpg"
    contact.save(contact_path, quality=85)
    with Image.open(contact_path) as check:
        if check.size != (800, 500) or widths != [400] * 4:
            raise RuntimeError("image validation failed")
    return {"images_processed": 4, "contact_bytes": contact_path.stat().st_size}


def search_files(directory):
    corpus = directory / "corpus"
    corpus.mkdir()
    expected = 0
    for index in range(1200):
        matched = index % 17 == 0
        expected += matched
        words = "status ordinary note " * 60
        if matched:
            words += "NEEDLE-ERROR-42 "
        (corpus / f"note-{index:04}.txt").write_text(words, encoding="utf-8")
    found = []
    digest = hashlib.sha256()
    for path in sorted(corpus.glob("*.txt")):
        content = path.read_bytes()
        digest.update(content)
        if b"NEEDLE-ERROR-42" in content:
            found.append(path.name)
    if len(found) != expected:
        raise RuntimeError("search result validation failed")
    return {"files_scanned": 1200, "matches": len(found), "sha256": digest.hexdigest()}


def run(workload, directory):
    if workload == "documents":
        return make_documents(directory)
    if workload == "data":
        return make_data(directory)
    if workload == "code":
        return run_code(directory)
    if workload == "images":
        return make_images(directory)
    if workload == "search":
        return search_files(directory)
    if workload == "workflow":
        data = make_data(directory, rows=20000)
        document = make_documents(directory, chart=directory / "chart.png")
        code = run_code(directory)
        return {"data": data, "documents": document, "code": code}
    raise ValueError("unknown workload")


if __name__ == "__main__":
    import resource
    workload, directory_name = sys.argv[1:3]
    directory = Path(directory_name)
    if workload not in WORKLOADS or not directory.is_dir() or directory.is_symlink():
        raise SystemExit("invalid workload or task directory")
    started = time.monotonic()
    details = run(workload, directory)
    (directory / "result.json").write_text(json.dumps({
        "workload": workload, "details": details,
        "elapsed_seconds": round(time.monotonic() - started, 3),
        "child_peak_rss_kb": resource.getrusage(resource.RUSAGE_CHILDREN).ru_maxrss,
    }), encoding="utf-8")
